"""Generates synthetic CV fixtures for pipeline testing (ADQ-3/5/8).

ST2i does not allow using real employee CVs, so this produces a fabricated
candidate pool instead — profiles aligned with the expert positions
requested in the real TDR samples in docs/tdrs/. The first 6 cover classic
IT-staffing profiles (project/team lead, software architecture,
database/BI, functional analysis, QA, DevOps). The next 5 (added once a
batch of AI/RAG/security/mobile/.NET TDRs were sourced) close the gaps
those expose: an AI/RAG engineer, a mobile + conversational-AI developer,
a cybersecurity/secure-development specialist, a .NET/C# enterprise
developer, and a software architect (microservices/design patterns/CI
toolchain). Mixed PDF/docx to exercise both parsers.

Run from backend/: `uv run python scripts/generate_synthetic_cvs.py`
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

OUTPUT_DIR = Path(__file__).resolve().parents[2] / "docs" / "cvs" / "synthetic"


@dataclass
class Experience:
    title: str
    employer: str
    sector: str
    start: str
    end: str | None
    responsibilities: list[str]


@dataclass
class Education:
    degree: str
    field: str
    institution: str
    year: int


@dataclass
class Certification:
    name: str
    issuing_body: str
    issue_date: str


@dataclass
class Candidate:
    full_name: str
    email: str
    phone: str
    birth_date: date
    summary: str
    education: list[Education]
    certifications: list[Certification]
    experience: list[Experience]
    languages: list[str]
    skills: list[str]
    file_format: str  # "pdf" | "docx"
    filename_stem: str = field(init=False)

    def __post_init__(self) -> None:
        self.filename_stem = self.full_name.lower().replace(" ", "-")


CANDIDATES: list[Candidate] = [
    Candidate(
        full_name="Nadia Ben Salah",
        email="nadia.bensalah@example.com",
        phone="+216 20 123 456",
        birth_date=date(1984, 4, 12),
        summary=(
            "Cheffe de projet SI avec 15 ans d'expérience dans le pilotage de "
            "programmes de transformation digitale pour le secteur public "
            "tunisien et international (Banque Mondiale, GIZ)."
        ),
        education=[
            Education("Diplôme d'ingénieur", "Génie Logiciel", "INSAT, Tunis", 2008),
            Education("Master", "Management de Projets IT", "IHEC Carthage", 2011),
        ],
        certifications=[
            Certification("PMP", "PMI", "2015-03"),
            Certification("PRINCE2 Practitioner", "AXELOS", "2017-06"),
        ],
        experience=[
            Experience(
                "Cheffe de projet SI",
                "ST2i",
                "Secteur public / GovTech",
                "2016-01",
                None,
                [
                    "Pilotage de la mise en place d'un système d'information de gestion "
                    "de la scolarité pour un ministère de l'éducation",
                    "Coordination d'une équipe pluridisciplinaire de 12 personnes "
                    "(développeurs, analystes fonctionnels, experts BI)",
                    "Interface avec les bailleurs de fonds internationaux (Banque "
                    "Mondiale, GIZ) et reporting d'avancement",
                ],
            ),
            Experience(
                "Consultante SI senior",
                "Deloitte Tunisie",
                "Conseil / secteur public",
                "2011-09",
                "2015-12",
                [
                    "Audit et cadrage de systèmes d'aide à la décision budgétaire",
                    "Rédaction de cahiers des charges pour des appels d'offres "
                    "internationaux",
                ],
            ),
        ],
        languages=["Arabe (natif)", "Français (courant)", "Anglais (B2)"],
        skills=[
            "Gestion de projet",
            "Méthodologies Agile/Scrum",
            "Systèmes d'information sectoriels",
            "Relations bailleurs de fonds",
            "Business Intelligence",
        ],
        file_format="pdf",
    ),
    Candidate(
        full_name="Karim Trabelsi",
        email="karim.trabelsi@example.com",
        phone="+216 22 987 654",
        birth_date=date(1989, 7, 3),
        summary=(
            "Architecte logiciel avec 10 ans d'expérience en conception de "
            "systèmes distribués et développement backend pour des projets "
            "gouvernementaux et bancaires."
        ),
        education=[
            Education("Diplôme d'ingénieur", "Informatique", "ENSI, Manouba", 2013),
        ],
        certifications=[
            Certification("AWS Certified Solutions Architect", "Amazon Web Services", "2020-04"),
        ],
        experience=[
            Experience(
                "Architecte logiciel",
                "ST2i",
                "Secteur public / IT",
                "2019-02",
                None,
                [
                    "Conception de l'architecture technique d'un applicatif multi-modules "
                    "(observation de classes, suivi de formation, distribution de manuels)",
                    "Mise en place de pipelines CI/CD et de bonnes pratiques DevOps",
                    "Encadrement technique d'une équipe de 6 développeurs backend/frontend",
                ],
            ),
            Experience(
                "Développeur backend senior",
                "Vermeg",
                "Fintech",
                "2013-09",
                "2019-01",
                [
                    "Développement de modules cœur bancaire en Java/Spring",
                    "Optimisation de requêtes PostgreSQL sur des volumes de données "
                    "importants",
                ],
            ),
        ],
        languages=["Arabe (natif)", "Français (courant)", "Anglais (courant)"],
        skills=[
            "Architecture logicielle",
            "Python",
            "Java / Spring",
            "PostgreSQL",
            "Docker & Kubernetes",
            "CI/CD",
        ],
        file_format="docx",
    ),
    Candidate(
        full_name="Salma Gharbi",
        email="salma.gharbi@example.com",
        phone="+216 24 555 321",
        birth_date=date(1992, 11, 20),
        summary=(
            "Experte base de données et Business Intelligence, 8 ans "
            "d'expérience dans la conception d'entrepôts de données et de "
            "tableaux de bord décisionnels pour le secteur public."
        ),
        education=[
            Education("Master", "Systèmes d'Information Décisionnels", "FSEG Tunis", 2016),
        ],
        certifications=[
            Certification("Oracle Database Administrator", "Oracle", "2018-05"),
            Certification("Microsoft Power BI Data Analyst", "Microsoft", "2021-11"),
        ],
        experience=[
            Experience(
                "Experte BI / Data",
                "ST2i",
                "Secteur public",
                "2020-06",
                None,
                [
                    "Conception d'un Delivery Dashboard pour le suivi de projets "
                    "d'investissement public",
                    "Modélisation de bases de données décisionnelles (schémas en étoile)",
                    "Développement de tableaux de bord Power BI pour des comités de "
                    "pilotage ministériels",
                ],
            ),
            Experience(
                "Administratrice de bases de données",
                "Tunisie Telecom",
                "Télécommunications",
                "2016-09",
                "2020-05",
                [
                    "Administration de bases Oracle et PostgreSQL en production",
                    "Mise en place de procédures de sauvegarde et de reprise après sinistre",
                ],
            ),
        ],
        languages=["Arabe (natif)", "Français (courant)", "Anglais (B1)"],
        skills=[
            "SQL / PL-SQL",
            "PostgreSQL",
            "Power BI",
            "Modélisation de données",
            "ETL",
        ],
        file_format="pdf",
    ),
    Candidate(
        full_name="Mehdi Aouadi",
        email="mehdi.aouadi@example.com",
        phone="+216 26 741 852",
        birth_date=date(1994, 2, 15),
        summary=(
            "Analyste fonctionnel avec 6 ans d'expérience dans le recueil de "
            "besoins et la rédaction de spécifications pour des projets "
            "sectoriels (éducation, santé)."
        ),
        education=[
            Education("Diplôme d'ingénieur", "Génie Logiciel", "ISIMM, Monastir", 2018),
        ],
        certifications=[
            Certification("Certified Business Analysis Professional", "IIBA", "2022-02"),
        ],
        experience=[
            Experience(
                "Analyste fonctionnel senior",
                "ST2i",
                "Secteur public / Éducation",
                "2021-03",
                None,
                [
                    "Recueil des besoins auprès d'un ministère de l'éducation pour un "
                    "système de gestion de la scolarité",
                    "Rédaction de spécifications fonctionnelles détaillées et de cas "
                    "d'usage",
                    "Animation d'ateliers utilisateurs et recette fonctionnelle",
                ],
            ),
            Experience(
                "Analyste fonctionnel",
                "Sofrecom Tunisie",
                "Télécommunications",
                "2018-09",
                "2021-02",
                [
                    "Analyse des processus métier et proposition d'optimisations SI",
                ],
            ),
        ],
        languages=["Arabe (natif)", "Français (courant)", "Anglais (B2)"],
        skills=[
            "Analyse fonctionnelle",
            "UML",
            "Recueil de besoins",
            "Rédaction de spécifications",
            "Tests de recette",
        ],
        file_format="docx",
    ),
    Candidate(
        full_name="Yosra Chaabane",
        email="yosra.chaabane@example.com",
        phone="+216 21 369 258",
        birth_date=date(1995, 9, 8),
        summary=(
            "Ingénieure QA avec 5 ans d'expérience en tests fonctionnels et "
            "automatisés pour des applications web du secteur public."
        ),
        education=[
            Education("Diplôme d'ingénieur", "Informatique", "ISI Ariana", 2019),
        ],
        certifications=[
            Certification("ISTQB Foundation Level", "ISTQB", "2020-09"),
        ],
        experience=[
            Experience(
                "Ingénieure QA",
                "ST2i",
                "Secteur public / IT",
                "2020-10",
                None,
                [
                    "Conception de plans de test pour un applicatif multi-modules "
                    "(observation de classes, suivi de formation)",
                    "Automatisation de tests de régression avec Selenium et Playwright",
                    "Suivi des anomalies et coordination avec les équipes de "
                    "développement",
                ],
            ),
        ],
        languages=["Arabe (natif)", "Français (courant)", "Anglais (B2)"],
        skills=[
            "Tests fonctionnels",
            "Tests automatisés",
            "Selenium",
            "Playwright",
            "Gestion des anomalies",
        ],
        file_format="pdf",
    ),
    Candidate(
        full_name="Amine Ferjani",
        email="amine.ferjani@example.com",
        phone="+216 23 147 963",
        birth_date=date(1993, 6, 25),
        summary=(
            "Expert DevOps et infrastructure, 7 ans d'expérience dans le "
            "déploiement et l'exploitation de plateformes cloud pour des "
            "projets institutionnels."
        ),
        education=[
            Education("Diplôme d'ingénieur", "Réseaux et Systèmes", "SUP'COM, Tunis", 2017),
        ],
        certifications=[
            Certification("Certified Kubernetes Administrator", "CNCF", "2021-07"),
            Certification("Terraform Associate", "HashiCorp", "2022-10"),
        ],
        experience=[
            Experience(
                "Expert DevOps",
                "ST2i",
                "Secteur public / IT",
                "2019-08",
                None,
                [
                    "Mise en place de l'infrastructure cloud pour un système "
                    "d'information ministériel (Docker, Kubernetes, CI/CD GitHub Actions)",
                    "Automatisation du provisioning d'infrastructure avec Terraform",
                    "Mise en place de supervision et d'alerting en production",
                ],
            ),
            Experience(
                "Administrateur systèmes",
                "Orange Tunisie",
                "Télécommunications",
                "2017-09",
                "2019-07",
                [
                    "Administration de serveurs Linux et supervision réseau",
                ],
            ),
        ],
        languages=["Arabe (natif)", "Français (courant)", "Anglais (courant)"],
        skills=[
            "Docker & Kubernetes",
            "Terraform",
            "CI/CD",
            "Linux",
            "Supervision (Prometheus/Grafana)",
        ],
        file_format="docx",
    ),
    # Added for ADQ-10+ taxonomy coverage: the original 6 candidates only
    # cover classic IT-staffing profiles (PM, backend, BI, functional
    # analysis, QA, DevOps infra). Real TDRs sourced since then (RAG/LLM
    # knowledge hubs, AI chatbots, secure web development, .NET/C#
    # enterprise systems, mobile+conversational-AI apps) expose real
    # candidate-pool gaps -- these 5 fill them, each grounded in an actual
    # skill/cert vocabulary pulled from those TDRs, not invented.
    Candidate(
        full_name="Ines Belhadj",
        email="ines.belhadj@example.com",
        phone="+216 25 741 369",
        birth_date=date(1994, 1, 30),
        summary=(
            "Ingénieure IA avec 6 ans d'expérience en traitement du langage "
            "naturel et architectures RAG, spécialisée dans la conception de "
            "hubs de connaissances et d'assistants conversationnels pour le "
            "secteur public."
        ),
        education=[
            Education("Diplôme d'ingénieur", "Data Science & Intelligence Artificielle", "ENSI, Manouba", 2018),
        ],
        certifications=[
            Certification(
                "Google Cloud Professional Machine Learning Engineer", "Google Cloud", "2023-05"
            ),
        ],
        experience=[
            Experience(
                "Ingénieure IA",
                "ST2i",
                "Secteur public / IA appliquée",
                "2022-03",
                None,
                [
                    "Conception d'une architecture RAG (Retrieval-Augmented Generation) pour "
                    "un hub de connaissances documentaire ministériel avec citations "
                    "vérifiables",
                    "Mise en place d'un pipeline d'ingestion et d'indexation vectorielle sur "
                    "une base de documents réglementaires",
                    "Développement d'un chatbot conversationnel basé sur des grands modèles "
                    "de langage (LLM), avec mitigation des hallucinations",
                ],
            ),
            Experience(
                "Data Scientist",
                "InstaDeep Tunis",
                "Tech / IA",
                "2018-09",
                "2022-02",
                [
                    "Traitement du langage naturel et analyse de sentiment sur des données "
                    "de réseaux sociaux",
                    "Développement de modèles de machine learning pour l'analyse prédictive",
                ],
            ),
        ],
        languages=["Arabe (natif)", "Français (courant)", "Anglais (courant)"],
        skills=[
            "Python",
            "Machine Learning",
            "Intelligence Artificielle Générative",
            "Grands modèles de langage",
            "Retrieval-Augmented Generation",
            "LangChain",
            "Traitement du langage naturel",
            "Prompt Engineering",
            "Base de données vectorielle",
        ],
        file_format="docx",
    ),
    Candidate(
        full_name="Youssef Kammoun",
        email="youssef.kammoun@example.com",
        phone="+216 27 852 963",
        birth_date=date(1995, 5, 17),
        summary=(
            "Développeur mobile avec 5 ans d'expérience dans la conception "
            "d'applications multi-plateformes et l'intégration d'assistants "
            "virtuels intelligents pour des projets à impact social."
        ),
        education=[
            Education("Diplôme d'ingénieur", "Génie Logiciel", "ISIMM, Monastir", 2019),
        ],
        certifications=[
            Certification("Associate Android Developer", "Google", "2021-08"),
        ],
        experience=[
            Experience(
                "Développeur Mobile Senior",
                "ST2i",
                "Secteur public / Inclusion sociale",
                "2021-06",
                None,
                [
                    "Développement d'une application mobile multi-plateforme (Flutter) pour "
                    "l'accompagnement de jeunes vers l'emploi et la formation",
                    "Intégration d'un assistant virtuel conversationnel (chatbot NLP) pour "
                    "guider le parcours utilisateur",
                    "Mise en place de l'authentification et d'une base de données temps réel "
                    "(Firebase)",
                ],
            ),
            Experience(
                "Développeur Full-Stack",
                "Expensya",
                "Tech / Fintech",
                "2019-02",
                "2021-05",
                [
                    "Développement backend Node.js et applications front-end React",
                    "Intégration d'API tierces et de services de paiement",
                ],
            ),
        ],
        languages=["Arabe (natif)", "Français (courant)", "Anglais (B2)"],
        skills=[
            "Flutter",
            "React",
            "Node.js",
            "Firebase",
            "MongoDB",
            "Chatbot conversationnel",
            "UX Design",
        ],
        file_format="pdf",
    ),
    Candidate(
        full_name="Rania Snoussi",
        email="rania.snoussi@example.com",
        phone="+216 28 963 741",
        birth_date=date(1991, 10, 4),
        summary=(
            "Experte cybersécurité avec 9 ans d'expérience en développement "
            "sécurisé et audit applicatif, spécialisée dans la remédiation "
            "des vulnérabilités web et la mise en place de mécanismes "
            "d'authentification unique pour des plateformes gouvernementales."
        ),
        education=[
            Education("Diplôme d'ingénieur", "Sécurité Informatique", "SUP'COM, Tunis", 2015),
        ],
        certifications=[
            Certification("CISSP", "(ISC)²", "2019-09"),
        ],
        experience=[
            Experience(
                "Experte Cybersécurité",
                "ST2i",
                "Secteur public / GovTech",
                "2020-01",
                None,
                [
                    "Audit de sécurité applicative (OWASP Top 10) sur des plateformes web "
                    "gouvernementales",
                    "Mise en œuvre de mécanismes d'authentification unique (SSO SAML / "
                    "OAuth 2.0)",
                    "Remédiation de vulnérabilités XSS, CSRF et injection SQL sur des "
                    "applications legacy",
                ],
            ),
            Experience(
                "Consultante Sécurité",
                "Deloitte Tunisie",
                "Conseil / Cybersécurité",
                "2015-09",
                "2019-12",
                [
                    "Tests d'intrusion et revues de code sécurisé pour des clients du "
                    "secteur financier",
                    "Sensibilisation des équipes de développement aux bonnes pratiques "
                    "OWASP",
                ],
            ),
        ],
        languages=["Arabe (natif)", "Français (courant)", "Anglais (courant)"],
        skills=[
            "Cybersécurité",
            "OWASP",
            "Injection SQL",
            "XSS",
            "CSRF",
            "SSO",
            "SAML",
            "OAuth 2.0",
            "Cryptographie applicative",
        ],
        file_format="docx",
    ),
    Candidate(
        full_name="Walid Mejri",
        email="walid.mejri@example.com",
        phone="+216 29 174 852",
        birth_date=date(1992, 8, 22),
        summary=(
            "Développeur .NET/C# senior avec 8 ans d'expérience dans la "
            "conception d'applications d'entreprise multi-niveaux pour la "
            "gestion financière publique."
        ),
        education=[
            Education("Diplôme d'ingénieur", "Informatique", "ENIT, Tunis", 2016),
        ],
        certifications=[
            Certification("Microsoft Azure Certified", "Microsoft", "2021-03"),
        ],
        experience=[
            Experience(
                "Développeur .NET Senior",
                "ST2i",
                "Secteur public / Systèmes d'information",
                "2019-09",
                None,
                [
                    "Développement d'applications ASP.NET MVC multi-niveaux pour la gestion "
                    "financière publique",
                    "Conception de modèles de données avec Entity Framework sur SQL Server",
                    "Migration d'applications legacy vers des environnements Azure",
                ],
            ),
            Experience(
                "Développeur C#",
                "Vermeg",
                "Fintech",
                "2016-09",
                "2019-08",
                [
                    "Développement de modules cœur métier en C#/.NET",
                    "Intégration de tableaux de bord Business Intelligence",
                ],
            ),
        ],
        languages=["Arabe (natif)", "Français (courant)", "Anglais (B2)"],
        skills=[
            ".NET",
            "ASP.NET MVC",
            "Entity Framework",
            "SQL Server",
            "Cloud computing",
            "Business Intelligence",
        ],
        file_format="pdf",
    ),
    Candidate(
        full_name="Nizar Chebbi",
        email="nizar.chebbi@example.com",
        phone="+216 20 486 357",
        birth_date=date(1989, 3, 11),
        summary=(
            "Architecte logiciel avec 11 ans d'expérience en conception "
            "d'architectures microservices et en mise en place de chaînes "
            "d'intégration continue pour des systèmes d'information "
            "sectoriels."
        ),
        education=[
            Education("Diplôme d'ingénieur", "Génie Logiciel", "INSAT, Tunis", 2013),
            Education("Master", "Architecture des Systèmes Logiciels", "Faculté des Sciences de Tunis", 2015),
        ],
        certifications=[
            Certification("TOGAF", "The Open Group", "2020-11"),
        ],
        experience=[
            Experience(
                "Architecte Logiciel",
                "ST2i",
                "Secteur public / IT",
                "2018-04",
                None,
                [
                    "Définition d'architectures microservices pour des systèmes "
                    "d'information multi-modules",
                    "Application de design patterns (GoF) pour la refactorisation "
                    "d'applications legacy",
                    "Mise en place de pipelines GitLab CI/CD avec analyse statique de code "
                    "(SonarQube) et formation des équipes au Test-Driven Development",
                ],
            ),
            Experience(
                "Ingénieur Logiciel Senior",
                "Sofrecom Tunisie",
                "Télécommunications",
                "2013-09",
                "2018-03",
                [
                    "Modélisation UML et développement d'applications Java/Spring",
                ],
            ),
        ],
        languages=["Arabe (natif)", "Français (courant)", "Anglais (courant)"],
        skills=[
            "Architecture logicielle",
            "Microservices",
            "Design Patterns",
            "MVC",
            "UML",
            "GitLab CI/CD",
            "SonarQube",
            "Test-Driven Development",
        ],
        file_format="docx",
    ),
]


def _format_experience_dates(exp: Experience) -> str:
    return f"{exp.start} – {exp.end or 'présent'}"


def render_pdf(candidate: Candidate, output_path: Path) -> None:
    styles = getSampleStyleSheet()
    name_style = ParagraphStyle(
        "Name", parent=styles["Title"], fontSize=18, spaceAfter=4
    )
    heading_style = ParagraphStyle(
        "SectionHeading", parent=styles["Heading2"], spaceBefore=12, spaceAfter=4
    )
    body_style = styles["BodyText"]

    doc = SimpleDocTemplate(
        str(output_path), pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm
    )
    story = [
        Paragraph(candidate.full_name, name_style),
        Paragraph(f"{candidate.email} — {candidate.phone}", body_style),
        Paragraph(f"Né(e) le: {candidate.birth_date:%d/%m/%Y}", body_style),
        Spacer(1, 8),
        Paragraph(candidate.summary, body_style),
        Paragraph("Formation", heading_style),
    ]
    for edu in candidate.education:
        story.append(
            Paragraph(
                f"{edu.degree} en {edu.field} — {edu.institution} ({edu.year})", body_style
            )
        )

    story.append(Paragraph("Certifications", heading_style))
    for cert in candidate.certifications:
        story.append(
            Paragraph(f"{cert.name} — {cert.issuing_body} ({cert.issue_date})", body_style)
        )

    story.append(Paragraph("Expérience professionnelle", heading_style))
    for exp in candidate.experience:
        story.append(
            Paragraph(
                f"<b>{exp.title}</b> — {exp.employer} ({exp.sector}), "
                f"{_format_experience_dates(exp)}",
                body_style,
            )
        )
        for resp in exp.responsibilities:
            story.append(Paragraph(f"• {resp}", body_style))

    story.append(Paragraph("Langues", heading_style))
    story.append(Paragraph(", ".join(candidate.languages), body_style))

    story.append(Paragraph("Compétences", heading_style))
    story.append(Paragraph(", ".join(candidate.skills), body_style))

    doc.build(story)


def render_docx(candidate: Candidate, output_path: Path) -> None:
    document = Document()
    document.add_heading(candidate.full_name, level=0)
    document.add_paragraph(f"{candidate.email} — {candidate.phone}")
    document.add_paragraph(f"Né(e) le: {candidate.birth_date:%d/%m/%Y}")
    document.add_paragraph(candidate.summary)

    document.add_heading("Formation", level=1)
    for edu in candidate.education:
        document.add_paragraph(
            f"{edu.degree} en {edu.field} — {edu.institution} ({edu.year})"
        )

    document.add_heading("Certifications", level=1)
    for cert in candidate.certifications:
        document.add_paragraph(f"{cert.name} — {cert.issuing_body} ({cert.issue_date})")

    document.add_heading("Expérience professionnelle", level=1)
    for exp in candidate.experience:
        document.add_paragraph(
            f"{exp.title} — {exp.employer} ({exp.sector}), {_format_experience_dates(exp)}"
        )
        for resp in exp.responsibilities:
            document.add_paragraph(resp, style="List Bullet")

    document.add_heading("Langues", level=1)
    document.add_paragraph(", ".join(candidate.languages))

    document.add_heading("Compétences", level=1)
    document.add_paragraph(", ".join(candidate.skills))

    document.save(str(output_path))


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for candidate in CANDIDATES:
        output_path = OUTPUT_DIR / f"{candidate.filename_stem}.{candidate.file_format}"
        if candidate.file_format == "pdf":
            render_pdf(candidate, output_path)
        else:
            render_docx(candidate, output_path)
        print(f"wrote {output_path.relative_to(OUTPUT_DIR.parents[2])}")


if __name__ == "__main__":
    main()
